"""
core/writers/docx - Utility for exporting OCR text to Word Documents (.docx).
"""

from __future__ import annotations

import io
import re
from typing import TYPE_CHECKING, Any

from docx import Document
from docx.shared import Inches, Pt

from omniscribe.core.writers.exporter_base import BaseDocumentExporter

if TYPE_CHECKING:
    from docx.document import Document as DocumentType
    from docx.text.paragraph import Paragraph

    from omniscribe.core.block_tree import DocumentTree
    from omniscribe.core.document import DocumentResult


class DocxMarkdownExporter(BaseDocumentExporter):
    """Document exporter producing Word (.docx) documents from Markdown or DocumentResult."""

    def export_document(self, document: DocumentResult, **kwargs: Any) -> io.BytesIO:
        """Render a DocumentResult to a Word document stream via markdown text."""
        raw_text = document.text() if callable(document.text) else str(document.text)
        return convert_markdown_to_docx(raw_text)

    def export_tree(self, tree: DocumentTree, **kwargs: Any) -> io.BytesIO:
        """Render a DocumentTree to a Word document stream."""
        from omniscribe.core.writers.docx_tree import convert_tree_to_docx

        return convert_tree_to_docx(tree)


def convert_markdown_to_docx(markdown_text: str) -> io.BytesIO:
    """
    Parse a Markdown string and generate a polished Word Document (.docx) as a binary stream.

    Supports:
    - Page breaks and titles (from '## Page X' or '--- PAGE X ---')
    - Markdown headings: # (H1), ## (H2), ### (H3)
    - Bullet lists (- or *)
    - Numbered lists (1. , etc.)
    - Inline styles: bold (**text**), italic (*text*), bold+italic (***text***), code (`text`)
    - Paragraph spacing and font customisation (Arial, 11pt, 1.15 line spacing)
    """
    doc = Document()

    # Configure page setup: Standard 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Configure base style: Arial 11pt
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(11)

    # Simple Markdown Parser
    # Guard: empty / whitespace-only input must not produce a content-free
    # DOCX (the file looks valid in Word but renders as a blank page, which
    # is indistinguishable from an export that "lost" the OCR text). Emit a
    # single visible placeholder paragraph so the user gets clear feedback
    # that the upstream content was empty instead of silently dropping it.
    if not markdown_text or not markdown_text.strip():
        placeholder = doc.add_paragraph()
        placeholder.paragraph_format.space_after = Pt(6)
        placeholder.paragraph_format.line_spacing = 1.15
        run = placeholder.add_run(
            "(No text recognized — run OCR processing to populate "
            "document content before exporting.)"
        )
        run.italic = True
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream

    lines = markdown_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Page Breaks (e.g., "## Page X" or "--- PAGE X ---")
        if stripped.startswith("## Page ") or (
            stripped.startswith("--- PAGE ") and stripped.endswith(" ---")
        ):
            _render_page_header(doc, stripped)
            i += 1
        elif _render_heading(doc, stripped) or _render_list_item(doc, stripped):
            i += 1
        elif stripped.startswith("|") and stripped.endswith("|"):
            i = _render_markdown_table(doc, lines, i)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            _add_inline_formatting(p, stripped)
            i += 1

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


def _render_page_header(doc: DocumentType, stripped: str) -> None:
    if doc.paragraphs:  # Skip page break for the very first page
        doc.add_page_break()
    heading_text = stripped.replace("-", "").replace("#", "").strip()
    h = doc.add_heading(heading_text, level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)


def _render_heading(doc: DocumentType, stripped: str) -> bool:
    if stripped.startswith("# "):
        h = doc.add_heading(stripped[2:], level=1)
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        return True
    if stripped.startswith("## "):
        h = doc.add_heading(stripped[3:], level=2)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
        return True
    if stripped.startswith("### "):
        h = doc.add_heading(stripped[4:], level=3)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(4)
        return True
    return False


def _render_list_item(doc: DocumentType, stripped: str) -> bool:
    if stripped.startswith("- ") or stripped.startswith("* "):
        content = stripped[2:]
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        _add_inline_formatting(p, content)
        return True
    num_match = re.match(r"^(\d+)\.\s+(.*)", stripped)
    if num_match:
        content = num_match.group(2)
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        _add_inline_formatting(p, content)
        return True
    return False


def _render_markdown_table(doc: DocumentType, lines: list[str], start_idx: int) -> int:
    idx = start_idx
    table_lines: list[str] = []
    while (
        idx < len(lines)
        and lines[idx].strip().startswith("|")
        and lines[idx].strip().endswith("|")
    ):
        table_lines.append(lines[idx].strip())
        idx += 1
    if len(table_lines) >= 2 and re.match(r"^\|[\s\-:|]+\|$", table_lines[1]):
        raw_headers = [c.strip() for c in table_lines[0][1:-1].split("|")]
        headers = raw_headers if raw_headers else [""]
        data_rows: list[list[str]] = []
        for row_line in table_lines[2:]:
            cells = [c.strip() for c in row_line[1:-1].split("|")]
            data_rows.append(cells)
        table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
        table.style = "Table Grid"
        for c_idx, h_text in enumerate(headers):
            hdr_cell = table.cell(0, c_idx)
            hdr_cell.text = h_text
            for p in hdr_cell.paragraphs:
                for r in p.runs:
                    r.bold = True
        for r_idx, row_cells in enumerate(data_rows):
            for c_idx in range(len(headers)):
                cell_val = row_cells[c_idx] if c_idx < len(row_cells) else ""
                table.cell(r_idx + 1, c_idx).text = cell_val
    else:
        for tl in table_lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            _add_inline_formatting(p, tl)
    return idx


def _add_inline_formatting(paragraph: Paragraph, text: str) -> None:
    """Helper to parse markdown bold/italic/code inline constructs and add runs to paragraph."""
    # Split text into formatted chunks and normal text
    pattern = re.compile(r"(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)")
    parts = pattern.split(text)

    for part in parts:
        if not part:
            continue

        if part.startswith("***") and part.endswith("***"):
            run = paragraph.add_run(part[3:-3])
            run.bold = True
            run.italic = True
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            paragraph.add_run(part)


__all__ = [
    "DocxMarkdownExporter",
    "convert_markdown_to_docx",
]
