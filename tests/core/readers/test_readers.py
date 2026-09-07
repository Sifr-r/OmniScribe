"""Unit tests for digital document readers (RFC 004 R2 Fast Path)."""

from __future__ import annotations

import io
from pathlib import Path

import pymupdf as fitz
import pytest
from docx import Document

from omniscribe.core.readers import (
    DocxReader,
    HtmlReader,
    MalformedDocumentError,
    MarkdownReader,
    ReaderError,
    get_reader_for_suffix,
    render_synthetic_pdf,
    supported_suffixes,
)

# -- DocxReader tests ---------------------------------------------------------


def test_docx_reader_success() -> None:
    doc = Document()
    doc.add_heading("Main Document Title", level=1)
    p1 = doc.add_paragraph("First introductory paragraph.")
    r1 = p1.add_run(" Bold text.")
    r1.bold = True
    r2 = p1.add_run(" Italic text.")
    r2.italic = True

    doc.add_paragraph("Bullet item 1", style="List Bullet")
    doc.add_paragraph("Bullet item 2", style="List Bullet")

    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Col A"
    tbl.cell(0, 1).text = "Col B"
    tbl.cell(1, 0).text = "Data A"
    tbl.cell(1, 1).text = "Data B"

    doc.add_heading("Sub Section Heading", level=2)
    doc.add_paragraph("Paragraph under subsection.")

    # Next section starts on page 2
    doc.add_heading("Chapter 2", level=1)
    doc.add_paragraph("Second page text.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    reader = DocxReader()
    result = reader.read(docx_bytes, filename="test.docx")

    assert result.source_path == "test.docx"
    assert len(result.pages) == 2
    assert result.tree is not None
    assert len(result.tree.pages) == 2
    assert len(result.tree.sections) == 3
    assert result.tree.sections[0].title == "Main Document Title"
    assert result.tree.sections[0].level == 1
    assert result.tree.sections[1].title == "Sub Section Heading"
    assert result.tree.sections[1].level == 2
    assert result.tree.sections[2].title == "Chapter 2"
    assert result.tree.sections[2].level == 1
    assert len(result.tree.tables) == 1

    # Check trust flags and confidence on every block
    for page in result.pages:
        for block in page.blocks:
            assert block.confidence == 1.0
            assert block.trust_score == 1.0
            assert block.trust_flags == ("source:digital",)
            assert block.source_processor == "digital"
            assert 0.0 <= block.bbox[0] <= 1.0
            assert 0.0 <= block.bbox[1] <= 1.0
            assert 0.0 <= block.bbox[2] <= 1.0
            assert 0.0 <= block.bbox[3] <= 1.0
            assert block.bbox[2] > block.bbox[0]
            assert block.bbox[3] > block.bbox[1]


def test_docx_reader_error_handling(tmp_path: Path) -> None:
    reader = DocxReader()

    # Empty bytes
    with pytest.raises(MalformedDocumentError):
        reader.read(b"", filename="empty.docx")

    # Corrupted / invalid docx
    with pytest.raises(MalformedDocumentError):
        reader.read(b"not a valid zip file", filename="corrupt.docx")

    # Non-existent file
    missing = tmp_path / "non_existent.docx"
    with pytest.raises(ReaderError):
        reader.read(missing)


# -- HtmlReader tests ---------------------------------------------------------


def test_html_reader_success() -> None:
    html = """<!DOCTYPE html>
<html>
<head><title>Test Document</title></head>
<body>
<h1>Page 1: Title</h1>
<p>Hello world paragraph with &amp; entities.</p>
<ul>
    <li>Bullet item 1</li>
    <li>Bullet item 2</li>
</ul>
<pre><code>print("Hello from code")</code></pre>
<table>
    <tr><th>Head 1</th><th>Head 2</th></tr>
    <tr><td>Cell 1</td><td>Cell 2</td></tr>
</table>
<hr/>
<h1>Page 2: Second Chapter</h1>
<p>Content for the second page.</p>
</body>
</html>"""

    reader = HtmlReader()
    result = reader.read(html.encode("utf-8"), filename="test.html")

    assert result.source_path == "test.html"
    assert len(result.pages) == 2
    assert result.tree is not None
    assert len(result.tree.pages) == 2
    assert len(result.tree.sections) == 2
    assert len(result.tree.tables) == 1

    # Verify blocks
    p0_kinds = [b.kind for b in result.pages[0].blocks]
    assert "section_header" in p0_kinds
    assert "paragraph" in p0_kinds
    assert "list_item" in p0_kinds
    assert "code" in p0_kinds
    assert "table" in p0_kinds

    for page in result.pages:
        for block in page.blocks:
            assert block.confidence == 1.0
            assert block.trust_score == 1.0
            assert block.trust_flags == ("source:digital",)
            assert block.source_processor == "digital"


def test_html_reader_error_handling(tmp_path: Path) -> None:
    reader = HtmlReader()

    with pytest.raises(MalformedDocumentError):
        reader.read(b"", filename="empty.html")

    missing = tmp_path / "missing.html"
    with pytest.raises(ReaderError):
        reader.read(missing)


# -- MarkdownReader tests -----------------------------------------------------


def test_markdown_reader_success() -> None:
    md = """# Title Heading 1

First introductory paragraph in markdown format.

## Sub Heading 2

- List item one
- List item two

~~~python
def calculate():
    return 100
~~~

> A simple blockquote line.

| Column Alpha | Column Beta |
| --- | --- |
| Row 1 Alpha | Row 1 Beta |
| Row 2 Alpha | Row 2 Beta |

---

# Page 2 Heading

This text is placed on page two.
"""

    reader = MarkdownReader()
    result = reader.read(md.encode("utf-8"), filename="test.md")

    assert result.source_path == "test.md"
    assert len(result.pages) == 2
    assert result.tree is not None
    assert len(result.tree.pages) == 2
    assert len(result.tree.sections) == 3
    assert result.tree.sections[0].title == "Title Heading 1"
    assert result.tree.sections[1].title == "Sub Heading 2"
    assert result.tree.sections[2].title == "Page 2 Heading"
    assert len(result.tree.tables) == 1

    p0_kinds = [b.kind for b in result.pages[0].blocks]
    assert "section_header" in p0_kinds
    assert "paragraph" in p0_kinds
    assert "list_item" in p0_kinds
    assert "code" in p0_kinds
    assert "table" in p0_kinds

    for page in result.pages:
        for block in page.blocks:
            assert block.confidence == 1.0
            assert block.trust_score == 1.0
            assert block.trust_flags == ("source:digital",)


def test_markdown_reader_error_handling(tmp_path: Path) -> None:
    reader = MarkdownReader()

    with pytest.raises(MalformedDocumentError):
        reader.read(b"", filename="empty.md")

    missing = tmp_path / "missing.md"
    with pytest.raises(ReaderError):
        reader.read(missing)


# -- Dispatch registry tests --------------------------------------------------


def test_dispatch_registry() -> None:
    assert isinstance(get_reader_for_suffix(".docx"), DocxReader)
    assert isinstance(get_reader_for_suffix("docx"), DocxReader)
    assert isinstance(get_reader_for_suffix(".html"), HtmlReader)
    assert isinstance(get_reader_for_suffix(".htm"), HtmlReader)
    assert isinstance(get_reader_for_suffix(".md"), MarkdownReader)
    assert isinstance(get_reader_for_suffix(".markdown"), MarkdownReader)

    # Unsupported
    assert get_reader_for_suffix(".pdf") is None
    assert get_reader_for_suffix(".png") is None
    assert get_reader_for_suffix(".xyz") is None
    assert get_reader_for_suffix("") is None

    # Supported suffixes list
    suffixes = supported_suffixes()
    assert ".docx" in suffixes
    assert ".html" in suffixes
    assert ".htm" in suffixes
    assert ".md" in suffixes
    assert ".markdown" in suffixes


# -- Synthetic PDF rendering tests --------------------------------------------


def test_render_synthetic_pdf() -> None:
    md = "# Overview\n\nDigital ingest produces a clean PDF.\n\n- Feature A\n- Feature B"
    reader = MarkdownReader()
    doc_result = reader.read(md.encode("utf-8"), filename="overview.md")

    pdf_bytes = render_synthetic_pdf(doc_result)
    assert len(pdf_bytes) > 0

    # Open with PyMuPDF and verify text layer
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    assert doc.page_count >= 1
    page_text = doc[0].get_text()
    assert "Overview" in page_text
    assert "Digital ingest" in page_text
    assert "Feature A" in page_text
    doc.close()
