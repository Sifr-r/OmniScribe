"""Integration tests for Digital-Document Ingest Fast Path (RFC 004 R2).

Verifies that uploads of DOCX, HTML, and Markdown files to /api/process completely
bypass OCRPipeline Surya and VLM inference, succeed with 200, emit trust summaries
with 1.0 confidence and source:digital flags, and preserve DOCX text parity.
"""

from __future__ import annotations

import asyncio
import io
import json
import time
from typing import Any

import httpx
import pymupdf as fitz
import pytest
from docx import Document
from fastapi import FastAPI

from omniscribe.core.readers import DocxReader
from omniscribe.core.writers.docx_tree import convert_tree_to_docx
from omniscribe.harness.context import Context
from omniscribe.plugins import artifacts as art
from omniscribe.plugins import jobs, progress, runtime
from omniscribe.plugins import state_backend as sb
from omniscribe.plugins.ocr.plugin import OCRPlugin


@pytest.fixture(autouse=True)
def forbid_vlm_and_surya_engines(monkeypatch: pytest.MonkeyPatch) -> None:
    """CRITICAL: Raise if HybridEngine or GroundedEngine is invoked.

    Digital-document fast path must completely bypass Surya and VLM inference.
    """
    def forbidden_hybrid(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError(
            "HybridEngine.execute was invoked! Digital documents must use reader fast path."
        )

    def forbidden_grounded(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError(
            "GroundedEngine.execute was invoked! Digital documents must use reader fast path."
        )

    monkeypatch.setattr(
        "omniscribe.core.workflows.hybrid.HybridEngine.execute",
        forbidden_hybrid,
    )
    monkeypatch.setattr(
        "omniscribe.core.workflows.grounded.GroundedEngine.execute",
        forbidden_grounded,
    )


async def _boot_app(**ocr_config: Any) -> tuple[Context, FastAPI]:
    ctx = Context()
    await ctx.plugin(runtime.RuntimePlugin(), config={})
    await ctx.plugin(sb.StateBackendPlugin(), config={"backend": "memory"})
    await ctx.plugin(art.ArtifactsPlugin(), config={})
    await ctx.plugin(jobs.JobsPlugin(), config={})
    await ctx.plugin(progress.ProgressPlugin(), config={})
    await ctx.plugin(OCRPlugin(), config=ocr_config)
    app = FastAPI()
    for router in ctx.routes():
        app.include_router(router)
    return ctx, app


def _client(app: FastAPI) -> httpx.AsyncClient:
    return httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    )


async def _wait_for_job(
    client: httpx.AsyncClient, job_id: str, target_status: str = "complete", timeout: float = 5.0
) -> dict[str, Any]:
    deadline = time.time() + timeout
    body: dict[str, Any] = {}
    while time.time() < deadline:
        res = await client.get(f"/api/process/status/{job_id}")
        body = res.json()
        if body.get("status") == target_status:
            return body
        if body.get("status") in ("failed", "cancelled", "error"):
            raise AssertionError(f"Job {job_id} terminated prematurely: {body}")
        await asyncio.sleep(0.05)
    raise AssertionError(f"Job {job_id} did not reach {target_status!r} within {timeout}s: {body}")


# -- Sync Fast Path Tests -----------------------------------------------------


async def test_docx_sync_upload_fast_path() -> None:
    # 1. Create a DOCX in memory
    doc = Document()
    doc.add_heading("Digital DOCX Ingest Test", level=1)
    doc.add_paragraph("This paragraph was authored in Word and ingested digitally.")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Col1"
    tbl.cell(0, 1).text = "Col2"
    tbl.cell(1, 0).text = "Alpha"
    tbl.cell(1, 1).text = "Beta"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    ctx, app = await _boot_app()
    try:
        async with _client(app) as client:
            response = await client.post(
                "/api/process",
                files={
                    "file": (
                        "test_input.docx",
                        docx_bytes,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )

        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/pdf")
        assert "x-text-artifact-id" in response.headers
        assert "x-text-artifact-token" in response.headers

        # Verify trust summary: 1.0 confidence and digital flag
        assert "x-document-trust" in response.headers
        trust_summary = json.loads(response.headers["x-document-trust"])
        assert trust_summary["average"] == 1.0
        assert trust_summary["flag_counts"].get("source:digital", 0) > 0

        # Verify PDF contents
        pdf_bytes = response.content
        fitz_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        assert fitz_doc.page_count >= 1
        page_text = fitz_doc[0].get_text()
        assert "Digital DOCX Ingest Test" in page_text
        assert "ingested digitally" in page_text
        fitz_doc.close()
    finally:
        await ctx.dispose()


async def test_html_sync_upload_fast_path() -> None:
    html_content = b"""<!DOCTYPE html>
<html>
<body>
<h1>Fast Path HTML Title</h1>
<p>Direct HTML ingest without any OCR inference.</p>
<table>
    <tr><th>Item</th><th>Price</th></tr>
    <tr><td>Widget</td><td>$10</td></tr>
</table>
</body>
</html>"""

    ctx, app = await _boot_app()
    try:
        async with _client(app) as client:
            response = await client.post(
                "/api/process",
                files={"file": ("page.html", html_content, "text/html")},
            )

        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/pdf")
        assert "x-document-trust" in response.headers
        trust = json.loads(response.headers["x-document-trust"])
        assert trust["average"] == 1.0
        assert trust["flag_counts"]["source:digital"] > 0

        pdf_doc = fitz.open(stream=response.content, filetype="pdf")
        text = pdf_doc[0].get_text()
        assert "Fast Path HTML Title" in text
        assert "Direct HTML ingest" in text
        pdf_doc.close()
    finally:
        await ctx.dispose()


async def test_markdown_sync_upload_fast_path() -> None:
    md_content = b"""# Markdown Ingest Benchmark

This file verifies MD ingest fast path.

- Item 1
- Item 2

| Key | Value |
| --- | --- |
| Foo | Bar |
"""

    ctx, app = await _boot_app()
    try:
        async with _client(app) as client:
            response = await client.post(
                "/api/process",
                files={"file": ("doc.md", md_content, "text/markdown")},
            )

        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/pdf")
        trust = json.loads(response.headers["x-document-trust"])
        assert trust["average"] == 1.0
        assert trust["flag_counts"]["source:digital"] > 0

        pdf_doc = fitz.open(stream=response.content, filetype="pdf")
        text = pdf_doc[0].get_text()
        assert "Markdown Ingest Benchmark" in text
        assert "Item 1" in text
        pdf_doc.close()
    finally:
        await ctx.dispose()


# -- Async Fast Path Tests ----------------------------------------------------


async def test_digital_async_submit_fast_path() -> None:
    md_content = b"# Async Heading\n\nTesting async job queue with digital reader."

    ctx, app = await _boot_app()
    try:
        async with _client(app) as client:
            res = await client.post(
                "/api/process/async",
                files={"file": ("notes.md", md_content, "text/markdown")},
            )
            assert res.status_code == 202
            job_id = res.json()["job_id"]

            status = await _wait_for_job(client, job_id, "complete")
            assert status["status"] == "complete"
            assert not status.get("error")
    finally:
        await ctx.dispose()


# -- DOCX In -> DOCX Out Text Parity Test -------------------------------------


def test_docx_in_to_docx_out_parity() -> None:
    """Verify text parity: DOCX input -> DocxReader -> DocumentTree -> convert_tree_to_docx."""
    orig_doc = Document()
    orig_doc.add_heading("Parity Test Heading 1", level=1)
    orig_doc.add_paragraph("First paragraph describing parity verification.")
    orig_doc.add_paragraph("Second bullet paragraph.", style="List Bullet")
    orig_doc.add_heading("Subheading Level 2", level=2)
    tbl = orig_doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Header A"
    tbl.cell(0, 1).text = "Header B"
    tbl.cell(1, 0).text = "Data 1"
    tbl.cell(1, 1).text = "Data 2"

    in_buf = io.BytesIO()
    orig_doc.save(in_buf)
    docx_in_bytes = in_buf.getvalue()

    # Step 1: Read via DocxReader
    reader = DocxReader()
    doc_result = reader.read(docx_in_bytes, filename="parity_input.docx")
    assert doc_result.tree is not None

    # Step 2: Export tree to DOCX via convert_tree_to_docx
    out_stream = convert_tree_to_docx(doc_result.tree)
    exported_docx_bytes = out_stream.getvalue()
    assert len(exported_docx_bytes) > 0

    # Step 3: Inspect exported DOCX and check text parity
    exported_doc = Document(io.BytesIO(exported_docx_bytes))
    exported_para_texts = [p.text.strip() for p in exported_doc.paragraphs if p.text.strip()]
    exported_table_texts: list[str] = []
    for table in exported_doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    exported_table_texts.append(text)

    # Verify input texts are faithfully represented in the exported DOCX
    assert "Parity Test Heading 1" in exported_para_texts
    assert "First paragraph describing parity verification." in exported_para_texts
    assert "Second bullet paragraph." in exported_para_texts
    assert "Subheading Level 2" in exported_para_texts
    assert "Header A" in exported_table_texts
    assert "Data 1" in exported_table_texts
