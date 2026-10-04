"""Rich document artifact: persistence, lossless round-trip, export fidelity.

Covers Finding 1 of ``docs/goal-alignment-assessment-2026-10-03.md``: the
``DocumentResult`` produced by the OCR processors (and by the digital
readers) must survive into a stored artifact and back out through
``build_tree`` / ``build_document_export`` instead of being flattened to
text lines. Also covers the server half of Finding 6 (per-page failures
reaching the status / list responses).
"""

from __future__ import annotations

import asyncio
import io
import json
import time
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import httpx
import pytest
from docx import Document as DocxDocument
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse

from omniscribe.core.block_tree import (
    BlockNode,
    BlockType,
    DocumentTree,
    PageTree,
    Section,
    TableNode,
)
from omniscribe.core.document import (
    DocumentBlock,
    DocumentPage,
    DocumentResult,
    DocumentSerializationError,
)
from omniscribe.core.ocr_quality import OCrQualitySettings, build_trust_orchestrator
from omniscribe.core.workflows.base import EngineBase
from omniscribe.harness.context import Context
from omniscribe.plugins import artifacts as art
from omniscribe.plugins import jobs, progress, runtime
from omniscribe.plugins import state_backend as sb
from omniscribe.plugins._http import plugin_error_exception_handler
from omniscribe.plugins.artifacts import ArtifactStore
from omniscribe.plugins.documents.artifact import (
    DOCUMENT_ARTIFACT_CONTENT_TYPE,
    DocumentArtifactError,
    decode_document_artifact,
    document_artifact_ref,
    encode_document_artifact,
    load_document_result,
)
from omniscribe.plugins.documents.routes import build_documents_router
from omniscribe.plugins.documents.service import (
    build_document_export,
    build_tree,
    load_pages,
)
from omniscribe.plugins.errors import PluginError
from omniscribe.plugins.ocr import service as ocr_service
from omniscribe.plugins.ocr.plugin import OCRPlugin

PDF_BYTES = b"%PDF-1.4 fake"

# -- fixtures ------------------------------------------------------------------


def _cell(
    text: str, page_idx: int, bbox: tuple[float, float, float, float]
) -> BlockNode:
    return BlockNode(
        block_type=BlockType.PARAGRAPH,
        bbox=bbox,
        text=text,
        page_idx=page_idx,
        confidence=0.9,
    )


def _table(page_idx: int) -> TableNode:
    return TableNode(
        rows=2,
        cols=2,
        page_idx=page_idx,
        bbox=(0.05, 0.20, 0.95, 0.32),
        cells=[
            [
                _cell("Item", page_idx, (0.05, 0.20, 0.45, 0.26)),
                _cell("Qty", page_idx, (0.45, 0.20, 0.95, 0.26)),
            ],
            [
                _cell("Widget", page_idx, (0.05, 0.26, 0.45, 0.32)),
                _cell("42", page_idx, (0.45, 0.26, 0.95, 0.32)),
            ],
        ],
    )


def _rich_document() -> DocumentResult:
    """A two-page result with a heading, a table, body text, and trust data."""
    table = _table(0)
    tree = DocumentTree(
        pages=[
            PageTree(
                page_idx=0,
                width=1700,
                height=2200,
                children=[
                    BlockNode(
                        block_type=BlockType.SECTION_HEADER,
                        bbox=(0.05, 0.05, 0.95, 0.10),
                        text="Quarterly Report",
                        page_idx=0,
                        block_id="b000000000000001",
                        level=1,
                        section_hierarchy=["Quarterly Report"],
                        confidence=0.98,
                        trust_score=0.97,
                        trust_flags=("HIGH_CONFIDENCE",),
                    ),
                    table,
                    BlockNode(
                        block_type=BlockType.PARAGRAPH,
                        bbox=(0.05, 0.36, 0.95, 0.50),
                        text="Revenue grew across every region this quarter.",
                        page_idx=0,
                        block_id="b000000000000002",
                        section_hierarchy=["Quarterly Report"],
                        confidence=0.55,
                        trust_score=0.42,
                        trust_flags=("LOW_CALIBRATED_CONF",),
                    ),
                ],
                metadata={"rotation": 0},
            ),
            PageTree(
                page_idx=1,
                width=1700,
                height=2200,
                children=[
                    BlockNode(
                        block_type=BlockType.SECTION_HEADER,
                        bbox=(0.05, 0.05, 0.95, 0.10),
                        text="Details",
                        page_idx=1,
                        block_id="b000000000000003",
                        level=2,
                        section_hierarchy=["Quarterly Report", "Details"],
                        confidence=0.93,
                        trust_score=0.9,
                    ),
                    BlockNode(
                        block_type=BlockType.PARAGRAPH,
                        bbox=(0.05, 0.12, 0.95, 0.30),
                        text="Appendix material follows on the second page.",
                        page_idx=1,
                        block_id="b000000000000004",
                        section_hierarchy=["Quarterly Report", "Details"],
                        confidence=0.88,
                    ),
                ],
            ),
        ],
        sections=[
            Section(title="Quarterly Report", level=1, start_page=0),
            Section(
                title="Details", level=2, start_page=1, block_id="b000000000000003"
            ),
        ],
        tables=[table],
    )
    return DocumentResult(
        pages=[
            DocumentPage(
                page_index=0,
                blocks=[
                    DocumentBlock(
                        bbox=(0.05, 0.05, 0.95, 0.10),
                        text="Quarterly Report",
                        kind="heading",
                        confidence=0.98,
                        source_processor="surya",
                        reading_order=0,
                        metadata={"label": "section_header", "font_size": 24},
                        trust_score=0.97,
                        trust_flags=("HIGH_CONFIDENCE",),
                    ),
                    DocumentBlock(
                        bbox=(0.05, 0.20, 0.95, 0.32),
                        text="Item Qty Widget 42",
                        kind="table",
                        confidence=0.9,
                        source_processor="structure",
                        reading_order=1,
                        metadata={"rows": 2, "cols": 2},
                    ),
                    DocumentBlock(
                        bbox=(0.05, 0.36, 0.95, 0.50),
                        text="Revenue grew across every region this quarter.",
                        kind="text",
                        confidence=0.55,
                        source_processor="surya",
                        reading_order=2,
                        metadata={"label": "paragraph"},
                        trust_score=0.42,
                        trust_flags=("LOW_CALIBRATED_CONF",),
                    ),
                ],
                width=1700,
                height=2200,
                metadata={"rotation": 0, "dpi": 200},
            ),
            DocumentPage(
                page_index=1,
                blocks=[
                    DocumentBlock(
                        bbox=(0.05, 0.05, 0.95, 0.10),
                        text="Details",
                        kind="heading",
                        confidence=0.93,
                        source_processor="surya",
                        reading_order=0,
                        metadata={"label": "section_header"},
                    ),
                    DocumentBlock(
                        bbox=(0.05, 0.12, 0.95, 0.30),
                        text="Appendix material follows on the second page.",
                        kind="text",
                        confidence=0.88,
                        source_processor="surya",
                        reading_order=1,
                    ),
                ],
                width=1700,
                height=2200,
            ),
        ],
        source_path="input.pdf",
        tree=tree,
    )


def _legacy_pages() -> dict[int, list[str]]:
    """Exactly what the legacy text artifact would hold for the same run."""
    return {
        0: [
            "Quarterly Report",
            "Item Qty Widget 42",
            "Revenue grew across every region this quarter.",
        ],
        1: ["Details", "Appendix material follows on the second page."],
    }


def _scored_pipeline(
    document: DocumentResult | None,
    *,
    failed_pages: list[int] | None = None,
) -> SimpleNamespace:
    return SimpleNamespace(
        last_document_result=document,
        last_failed_pages=failed_pages or [],
    )


def _patch_bridge(
    monkeypatch: pytest.MonkeyPatch,
    pipeline: Any,
    pages: dict[int, list[str]],
) -> None:
    def fake_build(settings: Any, request: Any, *, block_callbacks: Any = None) -> Any:
        return pipeline

    async def fake_run(
        pipeline_obj: Any,
        *,
        settings: Any,
        request: Any,
        input_path: str,
        output_path: str,
        on_progress: Any = None,
        on_warning: Any = None,
        cancel_check: Any = None,
    ) -> dict[int, list[str]]:
        Path(output_path).write_bytes(PDF_BYTES)
        return dict(pages)

    monkeypatch.setattr(ocr_service, "build_pipeline", fake_build)
    monkeypatch.setattr(ocr_service, "run_pipeline", fake_run)


async def _boot(**ocr_config: Any) -> tuple[Context, FastAPI]:
    ctx = Context()
    await ctx.plugin(runtime.RuntimePlugin(), config={})
    await ctx.plugin(sb.StateBackendPlugin(), config={"backend": "memory"})
    await ctx.plugin(art.ArtifactsPlugin(), config={})
    await ctx.plugin(jobs.JobsPlugin(), config={})
    await ctx.plugin(progress.ProgressPlugin(), config={})
    await ctx.plugin(OCRPlugin(), config=ocr_config)
    app = FastAPI()

    async def handle_plugin_error(request: Request, exc: Exception) -> JSONResponse:
        assert isinstance(exc, PluginError)
        return await plugin_error_exception_handler(request, exc)

    app.add_exception_handler(PluginError, handle_plugin_error)
    for router in ctx.routes():
        app.include_router(router)
    return ctx, app


def _client(app: FastAPI) -> httpx.AsyncClient:
    return httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    )


def _pdf_upload() -> dict[str, Any]:
    return {"files": {"file": ("a.pdf", b"%PDF-1.4 input", "application/pdf")}}


async def _wait_status(
    client: httpx.AsyncClient, job_id: str, status: str, *, timeout: float = 5.0
) -> dict[str, Any]:
    deadline = time.time() + timeout
    body: dict[str, Any] = {}
    while time.time() < deadline:
        body = (await client.get(f"/api/process/status/{job_id}")).json()
        if body.get("status") == status:
            return body
        await asyncio.sleep(0.01)
    raise AssertionError(f"job {job_id} never reached {status!r}; last={body}")


async def _fetch_artifact(
    ctx: Context, artifact_id: str, token: str
) -> tuple[bytes, str | None]:
    store: ArtifactStore = ctx.inject(ArtifactStore)
    blob = await store.get(artifact_id, token)
    assert blob is not None
    return blob.blob, blob.record.content_type


# -- core serialization --------------------------------------------------------


def test_document_result_roundtrip_is_lossless() -> None:
    original = _rich_document()
    restored = DocumentResult.from_dict(
        json.loads(json.dumps(original.to_dict(), ensure_ascii=False))
    )

    assert restored.source_path == "input.pdf"
    assert [p.page_index for p in restored.pages] == [0, 1]
    assert [(p.width, p.height) for p in restored.pages] == [(1700, 2200), (1700, 2200)]
    assert restored.pages[0].metadata == {"rotation": 0, "dpi": 200}

    heading = restored.pages[0].blocks[0]
    assert heading.bbox == (0.05, 0.05, 0.95, 0.10)
    assert heading.text == "Quarterly Report"
    assert heading.kind == "heading"
    assert heading.confidence == 0.98
    assert heading.source_processor == "surya"
    assert heading.reading_order == 0
    assert heading.metadata == {"label": "section_header", "font_size": 24}
    assert heading.trust_score == 0.97
    assert heading.trust_flags == ("HIGH_CONFIDENCE",)

    table_block = restored.pages[0].blocks[1]
    assert table_block.kind == "table"
    assert table_block.metadata == {"rows": 2, "cols": 2}
    assert table_block.trust_score is None

    assert restored.pages[0].blocks[2].trust_flags == ("LOW_CALIBRATED_CONF",)

    tree = restored.tree
    assert tree is not None
    assert len(tree.tables) == 1
    assert tree.tables[0].cells[1][0].text == "Widget"
    assert tree.tables[0].bbox == (0.05, 0.20, 0.95, 0.32)
    assert [(s.title, s.level, s.start_page) for s in tree.sections] == [
        ("Quarterly Report", 1, 0),
        ("Details", 2, 1),
    ]
    tree_heading = tree.pages[0].children[0]
    assert isinstance(tree_heading, BlockNode)
    assert tree_heading.section_hierarchy == ["Quarterly Report"]
    assert tree_heading.trust_flags == ("HIGH_CONFIDENCE",)
    assert tree_heading.block_id == "b000000000000001"
    assert tree.pages[0].metadata == {"rotation": 0}
    subheading = tree.pages[1].children[0]
    assert isinstance(subheading, BlockNode)
    assert subheading.section_hierarchy == [
        "Quarterly Report",
        "Details",
    ]


def test_document_artifact_envelope_roundtrips() -> None:
    blob = encode_document_artifact(_rich_document())
    envelope = json.loads(blob)
    assert envelope["schema"] == "omniscribe.document_result"
    assert envelope["version"] == 1
    restored = decode_document_artifact(blob)
    assert restored.tree is not None
    assert restored.tree.tables[0].cells[0][0].text == "Item"
    # Idempotent: encoding the restored result yields the same bytes.
    assert encode_document_artifact(restored) == blob


@pytest.mark.parametrize(
    "payload",
    [
        b"not json at all",
        json.dumps([1, 2, 3]).encode(),
        json.dumps({"schema": "something.else", "version": 1, "document": {}}).encode(),
        json.dumps({"schema": "omniscribe.document_result", "version": 99}).encode(),
        json.dumps(
            {
                "schema": "omniscribe.document_result",
                "version": 1,
                "document": {"pages": "not-a-list"},
            }
        ).encode(),
        json.dumps(
            {
                "schema": "omniscribe.document_result",
                "version": 1,
                "document": {
                    "pages": [
                        {
                            "page_index": 0,
                            "blocks": [{"bbox": [0.0, 0.0, 0.5], "text": "short bbox"}],
                        }
                    ]
                },
            }
        ).encode(),
        json.dumps(
            {
                "schema": "omniscribe.document_result",
                "version": 1,
                "document": {
                    "pages": [
                        {
                            "page_index": 0,
                            "blocks": [
                                {"bbox": [0.0, 0.0, 0.5, "x"], "text": "bad value"}
                            ],
                        }
                    ]
                },
            }
        ).encode(),
        # Valid pages, undecodable tree: must not degrade to an empty document.
        json.dumps(
            {
                "schema": "omniscribe.document_result",
                "version": 1,
                "document": {
                    "pages": [
                        {
                            "page_index": 0,
                            "blocks": [{"bbox": [0, 0, 1, 1], "text": "x"}],
                        }
                    ],
                    "tree": {
                        "pages": [
                            {
                                "page_idx": 0,
                                "children": [
                                    {
                                        "block_type": "not_a_real_block_type",
                                        "bbox": [0, 0, 1, 1],
                                        "text": "x",
                                        "page_idx": 0,
                                    }
                                ],
                            }
                        ]
                    },
                },
            }
        ).encode(),
    ],
)
def test_corrupt_document_artifact_raises_typed_error(payload: bytes) -> None:
    with pytest.raises(DocumentArtifactError) as excinfo:
        decode_document_artifact(payload)
    assert excinfo.value.error == "artifact_corrupt"
    assert excinfo.value.status_code == 500
    assert excinfo.value.detail


def test_truncated_tree_is_not_silently_dropped() -> None:
    """A broken tree must not yield a silently structure-less document."""
    payload = json.dumps(
        {
            "schema": "omniscribe.document_result",
            "version": 1,
            "document": {
                "pages": [
                    {"page_index": 0, "blocks": [{"bbox": [0, 0, 1, 1], "text": "x"}]}
                ],
                "tree": {
                    "pages": [
                        {
                            "page_idx": 0,
                            "children": [
                                {
                                    "block_type": "not_a_real_block_type",
                                    "bbox": [0, 0, 1, 1],
                                    "text": "x",
                                    "page_idx": 0,
                                }
                            ],
                        }
                    ]
                },
            },
        }
    ).encode()
    with pytest.raises(DocumentArtifactError) as excinfo:
        decode_document_artifact(payload)
    assert "tree" in excinfo.value.detail


def test_document_block_serialization_rejects_non_json_metadata() -> None:
    block = DocumentBlock(bbox=(0.0, 0.0, 1.0, 1.0), text="x", metadata={"bad": {1, 2}})
    with pytest.raises(DocumentSerializationError):
        DocumentBlock.to_dict(block)
    with pytest.raises(DocumentArtifactError):
        encode_document_artifact(
            DocumentResult(
                pages=[DocumentPage(page_index=0, blocks=[block])],
            )
        )


def test_document_artifact_ref_requires_both_halves() -> None:
    assert document_artifact_ref(
        {"document_artifact_id": "a", "document_artifact_token": "t"}
    ) == ("a", "t")
    assert document_artifact_ref({"document_artifact_id": "a"}) is None
    assert document_artifact_ref({"document_artifact_token": "t"}) is None
    assert document_artifact_ref({}) is None


# -- build_tree / exports ------------------------------------------------------


def test_build_tree_prefers_rich_document_artifact() -> None:
    document = decode_document_artifact(encode_document_artifact(_rich_document()))
    tree = build_tree(_legacy_pages(), document=document)

    header = tree.pages[0].children[0]
    assert isinstance(header, BlockNode)
    assert header.block_type == BlockType.SECTION_HEADER
    assert header.bbox == (0.05, 0.05, 0.95, 0.10)
    assert header.level == 1
    assert header.section_hierarchy == ["Quarterly Report"]
    assert header.trust_score == 0.97
    assert isinstance(tree.pages[0].children[1], TableNode)
    assert tree.tables[0].cells[1][1].text == "42"
    assert [s.title for s in tree.sections] == ["Quarterly Report", "Details"]


def test_build_tree_falls_back_to_text_artifact_only() -> None:
    """Old job / older client: only the legacy text artifact exists."""
    pages = load_pages({"0": "SHORT HEADING\nbody line one", "1": "second page"})
    tree = build_tree(pages)
    assert tree.pages[0].children[0].block_type == BlockType.SECTION_HEADER
    assert tree.pages[0].children[0].bbox == (0.0, 0.0, 0.0, 0.0)
    assert tree.pages[0].children[1].block_type == BlockType.PARAGRAPH
    assert tree.tables == []
    assert tree.sections == []


def test_build_document_export_carries_real_structure() -> None:
    document = decode_document_artifact(encode_document_artifact(_rich_document()))
    page_text = _legacy_pages()

    markdown = build_document_export(
        page_text=page_text, metadata=None, export_format="markdown", document=document
    )
    assert isinstance(markdown, str)
    assert "# Quarterly Report" in markdown
    assert "## Details" in markdown
    assert "| Item | Qty |" in markdown
    assert "| Widget | 42 |" in markdown

    for export_format in ("json", "docling", "mineru"):
        payload = build_document_export(
            page_text=page_text,
            metadata={"filename": "a.pdf"},
            export_format=export_format,
            document=document,
        )
        assert isinstance(payload, dict)
        structure = payload["structure"]
        assert structure["tables"][0]["cells"][1][0]["text"] == "Widget"
        assert structure["sections"][1]["title"] == "Details"
        assert structure["pages"][0]["children"][0]["bbox"] == [0.05, 0.05, 0.95, 0.10]
        assert structure["pages"][0]["children"][0]["trust_flags"] == [
            "HIGH_CONFIDENCE"
        ]
        # Legacy page-text keys are untouched (docling names it "document").
        legacy_pages = (
            payload["document"] if export_format == "docling" else payload["pages"]
        )
        assert legacy_pages[0]["page_index"] == 0

    plain = build_document_export(
        page_text=page_text, metadata=None, export_format="text", document=document
    )
    assert plain == build_document_export(
        page_text=page_text, metadata=None, export_format="text"
    )


# -- sync OCR path -------------------------------------------------------------


async def test_sync_ocr_persists_rich_artifact_and_legacy_text_artifact(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_bridge(monkeypatch, _scored_pipeline(_rich_document()), _legacy_pages())
    ctx, app = await _boot()
    try:
        async with _client(app) as client:
            response = await client.post("/api/process", **_pdf_upload())
        assert response.status_code == 200
        assert response.content == PDF_BYTES
        assert response.headers["x-text-artifact-id"]
        document_id = response.headers["x-document-artifact-id"]
        document_token = response.headers["x-document-artifact-token"]
        assert document_id != response.headers["x-text-artifact-id"]

        # The legacy text artifact keeps its exact shape.
        text_blob, _ = await _fetch_artifact(
            ctx,
            response.headers["x-text-artifact-id"],
            response.headers["x-text-artifact-token"],
        )
        assert json.loads(text_blob) == {
            str(page): "\n".join(lines) for page, lines in _legacy_pages().items()
        }

        rich_blob, content_type = await _fetch_artifact(
            ctx, document_id, document_token
        )
        assert content_type == DOCUMENT_ARTIFACT_CONTENT_TYPE
        document = decode_document_artifact(rich_blob)
        assert document.tree is not None
        assert document.pages[0].width == 1700
        assert document.pages[0].blocks[0].kind == "heading"
        assert document.pages[0].blocks[0].trust_flags == ("HIGH_CONFIDENCE",)

        # ... and it comes back out through the export builders.
        tree = build_tree(_legacy_pages(), document=document)
        assert tree.pages[0].children[0].bbox == (0.05, 0.05, 0.95, 0.10)
        markdown = build_document_export(
            page_text=_legacy_pages(),
            metadata=None,
            export_format="markdown",
            document=document,
        )
        assert "| Widget | 42 |" in markdown
    finally:
        await ctx.dispose()


async def test_sync_ocr_without_document_result_keeps_text_artifact_only(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A pipeline that recorded no DocumentResult must not invent one."""
    _patch_bridge(monkeypatch, _scored_pipeline(None), _legacy_pages())
    ctx, app = await _boot()
    try:
        async with _client(app) as client:
            response = await client.post("/api/process", **_pdf_upload())
        assert response.status_code == 200
        assert response.headers["x-text-artifact-id"]
        assert "x-document-artifact-id" not in response.headers
    finally:
        await ctx.dispose()


# -- async OCR path -----------------------------------------------------------


async def test_async_ocr_metadata_carries_document_handle(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_bridge(monkeypatch, _scored_pipeline(_rich_document()), _legacy_pages())
    ctx, app = await _boot()
    try:
        async with _client(app) as client:
            submit = await client.post("/api/process/async", **_pdf_upload())
            assert submit.status_code == 202
            job_id = submit.json()["job_id"]
            await _wait_status(client, job_id, "complete")

        record = await ctx.inject(jobs.JobQueue).status(job_id)
        assert record is not None
        ref = document_artifact_ref(record.request_meta)
        assert ref is not None, record.request_meta

        store: ArtifactStore = ctx.inject(ArtifactStore)
        document = await load_document_result(store, *ref)
        assert document is not None
        assert document.tree is not None
        assert document.tree.tables[0].cells[0][0].text == "Item"
        assert [p.page_index for p in document.pages] == [0, 1]
        assert document.pages[1].blocks[0].bbox == (0.05, 0.05, 0.95, 0.10)
        assert document.pages[0].metadata["dpi"] == 200

        # The text artifact is still exactly the legacy shape.
        text_ref = (
            record.request_meta["text_artifact_id"],
            record.request_meta["text_artifact_token"],
        )
        text_blob, _ = await _fetch_artifact(ctx, *text_ref)
        assert json.loads(text_blob) == {
            str(page): "\n".join(lines) for page, lines in _legacy_pages().items()
        }

        # Storing the same document twice keeps both artifacts independently
        # decodable (idempotent, no shared mutable state).
        second = await store.put(
            encode_document_artifact(document),
            content_type=DOCUMENT_ARTIFACT_CONTENT_TYPE,
            owner_job_id=job_id,
        )
        again = await load_document_result(store, second.id, second.token)
        assert again is not None
        assert again.to_dict() == document.to_dict()
    finally:
        await ctx.dispose()


async def test_load_document_result_returns_none_when_absent(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_bridge(monkeypatch, _scored_pipeline(_rich_document()), _legacy_pages())
    ctx, _app = await _boot()
    try:
        store: ArtifactStore = ctx.inject(ArtifactStore)
        assert await load_document_result(store, "missing", "tok") is None
    finally:
        await ctx.dispose()


# -- digital (DOCX) reader path ------------------------------------------------


def _docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Native Table Document", level=1)
    doc.add_paragraph("Body paragraph authored in Word.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Qty"
    table.cell(1, 0).text = "Widget"
    table.cell(1, 1).text = "42"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def test_digital_docx_ingest_retains_table_structure_through_export() -> None:
    docx_bytes = _docx_bytes()
    ctx, app = await _boot()
    try:
        async with _client(app) as client:
            response = await client.post(
                "/api/process",
                files={
                    "file": (
                        "native.docx",
                        docx_bytes,
                        "application/vnd.openxmlformats-officedocument"
                        ".wordprocessingml.document",
                    )
                },
            )
        assert response.status_code == 200
        rich_blob, _content_type = await _fetch_artifact(
            ctx,
            response.headers["x-document-artifact-id"],
            response.headers["x-document-artifact-token"],
        )
        document = decode_document_artifact(rich_blob)
        assert document.tree is not None
        assert len(document.tree.tables) == 1
        assert document.tree.tables[0].rows == 2
        assert document.tree.tables[0].cols == 2
        assert document.tree.tables[0].cells[1][0].text == "Widget"

        page_text = load_pages(
            {
                str(page.page_index): "\n".join(
                    b.text for b in page.blocks if b.text.strip()
                )
                for page in document.pages
            }
        )
        markdown = build_document_export(
            page_text=page_text,
            metadata=None,
            export_format="markdown",
            document=document,
        )
        assert "# Native Table Document" in markdown
        assert "| Widget | 42 |" in markdown
    finally:
        await ctx.dispose()


# -- failed pages (Finding 6, server half) -------------------------------------


async def test_recorded_page_failure_surfaces_in_status_and_list(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_bridge(
        monkeypatch,
        _scored_pipeline(_rich_document(), failed_pages=[2, 5]),
        _legacy_pages(),
    )
    ctx, app = await _boot()
    try:
        async with _client(app) as client:
            submit = await client.post("/api/process/async", **_pdf_upload())
            assert submit.status_code == 202
            job_id = submit.json()["job_id"]
            status = await _wait_status(client, job_id, "complete")
            listing = (await client.get("/api/jobs")).json()
            result = await client.get(
                f"/api/jobs/{job_id}/result",
                params={"token": submit.json()["result_token"]},
            )
        assert status["failed_pages"] == [2, 5]
        assert listing[0]["failed_pages"] == [2, 5]
        assert result.headers["x-failed-pages"] == "2,5"
    finally:
        await ctx.dispose()


async def test_failed_pages_default_to_empty_when_none_recorded(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_bridge(monkeypatch, _scored_pipeline(_rich_document()), _legacy_pages())
    ctx, app = await _boot()
    try:
        async with _client(app) as client:
            submit = await client.post("/api/process/async", **_pdf_upload())
            job_id = submit.json()["job_id"]
            status = await _wait_status(client, job_id, "complete")
            listing = (await client.get("/api/jobs")).json()
        assert status["failed_pages"] == []
        assert listing[0]["failed_pages"] == []
    finally:
        await ctx.dispose()


async def test_trust_survives_structure_artifact_and_http_export(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = _rich_document()
    document.pages[0].blocks[2].confidence = 0.2
    assert document.tree is not None
    original_tree = document.tree.to_dict()
    engine = EngineBase(
        output_writer=lambda *_args: None,
        trust_orchestrator=build_trust_orchestrator(
            OCrQualitySettings(calibration_enabled=True, trust_flag_threshold=0.7)
        ),
    )
    scored = await engine._apply_trust(document, model_id="unregistered-model")
    assert document.tree.to_dict() == original_tree
    assert scored.tree is not None
    assert [s.to_dict() for s in scored.tree.sections] == original_tree["sections"]
    subheading = scored.tree.pages[1].children[0]
    assert isinstance(subheading, BlockNode)
    assert subheading.level == 2
    assert subheading.section_hierarchy == [
        "Quarterly Report",
        "Details",
    ]
    body_node = scored.tree.pages[0].children[2]
    assert isinstance(body_node, BlockNode)
    assert body_node.trust_score == pytest.approx(0.2)
    assert body_node.trust_flags == ("low_calibrated_conf",)
    assert scored.tree.tables[0].cells[1][0].trust_score == pytest.approx(0.9)
    # Reapplying scoring changes neither structure nor the synchronized scores.
    again = await engine._apply_trust(scored, model_id="unregistered-model")
    assert again.to_dict() == scored.to_dict()
    _patch_bridge(monkeypatch, _scored_pipeline(scored), _legacy_pages())
    ctx, app = await _boot()
    app.include_router(build_documents_router(ctx))
    try:
        async with _client(app) as client:
            response = await client.post("/api/process", **_pdf_upload())
            assert response.status_code == 200
            body = {
                "text_artifact_id": response.headers["x-text-artifact-id"],
                "text_artifact_token": response.headers["x-text-artifact-token"],
                "document_artifact_id": response.headers["x-document-artifact-id"],
                "document_artifact_token": response.headers[
                    "x-document-artifact-token"
                ],
                "export_format": "json",
            }
            exported = await client.post("/api/export/document", json=body)
            assert exported.status_code == 200
        blob, _ = await _fetch_artifact(
            ctx, exported.json()["artifact_id"], exported.json()["token"]
        )
        structure = json.loads(blob)["structure"]
        assert structure["pages"][0]["children"][2]["trust_score"] == pytest.approx(0.2)
        assert structure["tables"][0]["cells"][1][0]["trust_score"] == pytest.approx(
            0.9
        )
        assert structure["sections"] == original_tree["sections"]
    finally:
        await ctx.dispose()


@pytest.mark.parametrize("async_mode", [False, True])
async def test_all_pages_failed_is_an_error_without_result_artifacts(
    monkeypatch: pytest.MonkeyPatch,
    async_mode: bool,
) -> None:
    _patch_bridge(
        monkeypatch,
        _scored_pipeline(_rich_document(), failed_pages=[0, 1]),
        {0: [], 1: []},
    )
    ctx, app = await _boot()
    try:
        async with _client(app) as client:
            path = "/api/process/async" if async_mode else "/api/process"
            response = await client.post(path, **_pdf_upload())
            if async_mode:
                assert response.status_code == 202
                status = await _wait_status(client, response.json()["job_id"], "error")
                assert "every processed page" in status["error"]
                assert status["text_artifact_id"] is None
                result = await client.get(
                    f"/api/jobs/{response.json()['job_id']}/result",
                    params={"token": response.json()["result_token"]},
                )
                assert result.status_code == 404
            else:
                assert response.status_code == 502
                assert response.json()["error"] == "ocr_all_pages_failed"
                assert "every processed page" in response.json()["detail"]
                assert "x-document-artifact-id" not in response.headers
    finally:
        await ctx.dispose()
